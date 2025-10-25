from mcp.server.fastmcp import FastMCP
from docx import Document
import json
import re
from gemini import ModifyWordDocument
from word import write_document
mcp = FastMCP('doc-agent')
@mcp.tool()
async def modify_text_of_doc(filename,topic):
    """
    This tool is used when the user wants only modify the contents of the the document so that aligns with a particular topic
    ARGS: filename,topic
    """
    res = []
    doc = Document(filename)
    for para in doc.paragraphs:
        
        res.append({"style":para.style.name,"content":para.text.strip()})
    
    with open("file_info.json",'w',encoding='utf-8') as f:
        json.dump(res,f,ensure_ascii=False)
    obj = ModifyWordDocument(res)
    result = obj.modify_document(topic)
    result = re.sub("\n","",result)
    print(result)
    with open("result.json",'w',encoding='utf-8') as f:
        json.dump(eval(result),f,ensure_ascii=False)
    

    res = write_document(filename,eval(result))
    if res == "Error":
        return "An error occurred"
    return "The action has been completed"
@mcp.tool()
async def modify_formatting_of_doc(filename,request):
    doc = Document(filename)
    res = []
    for para in doc.paragraphs:
        res.append({"style":para.style.name,"content":para.text.strip()})
    obj = ModifyWordDocument(res)
    res = eval(obj.modify_formatting(request))
    res = res['formatting']

    
    for para in doc.paragraphs:
        text = para.text

        for key in res.keys():
            if key == "bold":
                sentences = res[key]

                for sentence in sentences:
                    if sentence in text:
                        parts = text.split(sentence)

                        for run in para.runs():
                            run.clear()
                        
                        para._element.clear_content()

                        for i,part in enumerate(parts):
                            if part:
                                para.add_run(part)
                            
                            if i<len(parts)-1:
                                styled_sentence = para.add_run(sentence)
                                styled_sentence.bold = True
                                

mcp.run(transport='stdio')