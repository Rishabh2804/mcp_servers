from docx import Document

def write_document(path,content):
    title = content[0]['content']
    doc = Document()
    doc.add_paragraph(title,style='Title')

    for c in content:
        style = c['style']
        heading = style.split()
        try:
            
            if style == "Normal":
                doc.add_paragraph(c['content'])
                
           
            elif heading[0] == "Heading":

                doc.add_heading(c['content'],level=int(heading[1]))
        
        except:
            return "Error"
    
    doc.save(path)
    return "Success"